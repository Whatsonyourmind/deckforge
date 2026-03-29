"""Deal overview slide renderer -- one-pager with key metrics and traffic lights."""

from __future__ import annotations

import logging
from typing import TYPE_CHECKING

from deckforge.ir.elements.base import Position
from deckforge.rendering.slide_renderers.base import (
    BaseFinanceSlideRenderer,
    _infer_format,
)

if TYPE_CHECKING:
    from pptx.slide import Slide

    from deckforge.ir.slides.base import BaseSlide
    from deckforge.themes.types import ResolvedTheme

logger = logging.getLogger(__name__)

# Known traffic-light status keywords.
_TRAFFIC_STATUSES = frozenset({"green", "yellow", "red"})


class DealOverviewRenderer(BaseFinanceSlideRenderer):
    """Renders deal overview one-pager with metrics grid and traffic-light indicators."""

    def render(self, slide: Slide, ir_slide: BaseSlide, theme: ResolvedTheme) -> None:
        elements = ir_slide.elements

        # Title
        title = self._find_heading(elements) or "Deal Overview"
        title_pos = Position(x=0.75, y=0.4, width=11.8, height=0.8)
        self._add_title(slide, title, theme, title_pos)

        # Find all table elements
        tables = self._find_table_elements(elements)

        # Classify tables: metrics table vs status table
        metrics_table = None
        status_table = None

        for tbl in tables:
            content = tbl.content
            if self._is_status_table(content):
                status_table = tbl
            elif metrics_table is None:
                metrics_table = tbl

        # Render metrics table
        if metrics_table:
            content = metrics_table.content
            col_formats = [_infer_format(h) for h in content.headers]
            m_pos = Position(x=0.75, y=1.5, width=5.5, height=3.5)
            self._add_table(
                slide,
                headers=content.headers,
                rows=[list(r) for r in content.rows],
                theme=theme,
                position=m_pos,
                column_formats=col_formats,
                footer_row=list(content.footer_row) if content.footer_row else None,
            )

        # Render status indicators with traffic lights
        if status_table:
            content = status_table.content
            # Render as a table first
            s_pos = Position(x=7.0, y=1.5, width=4.5, height=2.5)
            self._add_table(
                slide,
                headers=content.headers,
                rows=[list(r) for r in content.rows],
                theme=theme,
                position=s_pos,
            )

            # Add traffic-light circles next to each status row
            for row_idx, row in enumerate(content.rows):
                # Look for a status value in the row
                for val in row:
                    if isinstance(val, str) and val.lower() in _TRAFFIC_STATUSES:
                        indicator_pos = Position(
                            x=11.8,
                            y=2.0 + row_idx * 0.6,
                            width=0.25,
                            height=0.25,
                        )
                        self._add_shape_indicator(slide, val, indicator_pos)
                        break

    @staticmethod
    def _is_status_table(content) -> bool:
        """Check if table contains traffic-light status values."""
        for row in content.rows:
            for val in row:
                if isinstance(val, str) and val.lower() in _TRAFFIC_STATUSES:
                    return True
        return False


__all__ = ["DealOverviewRenderer"]
